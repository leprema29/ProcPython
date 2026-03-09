import os
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.metier.variables import Variables


class BaseCsvToXlsx:
    """Base class for CSV to XLSX conversion"""

    def __init__(self, phone, date_requisition, demandeur_requisition=""):
        self.phone = phone
        self.date_requisition = date_requisition
        self.demandeur_requisition = demandeur_requisition
        self.base_dir = os.path.join(Variables.DESTINATION_DOSSIERS, date_requisition, phone)
        self.wb = Workbook()
        self.header_font = Font(bold=True, color="FFFFFF", size=10)
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

    def _read_csv(self, file_path):
        """Read CSV file and return rows"""
        rows = []
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                reader = csv.reader(f)
                for row in reader:
                    rows.append(row)
        except Exception as e:
            print(f"Error reading CSV {file_path}: {e}")
        return rows

    def _create_sheet(self, title, csv_file, headers=None):
        """Create a worksheet from CSV data"""
        csv_path = os.path.join(self.base_dir, csv_file)
        if not os.path.exists(csv_path):
            return None

        rows = self._read_csv(csv_path)
        if not rows:
            return None

        if self.wb.active and self.wb.active.title == 'Sheet':
            ws = self.wb.active
            ws.title = title
        else:
            ws = self.wb.create_sheet(title=title)

        # Use provided headers or first row
        if headers:
            ws.append(headers)
            start_row = 0
        else:
            ws.append(rows[0])
            start_row = 1

        # Style header row
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = self.thin_border

        # Add data rows
        for row_data in rows[start_row:]:
            ws.append(row_data)

        # Style data cells
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                cell.border = self.thin_border
                cell.alignment = Alignment(horizontal='center')

        # Auto-adjust column widths
        for col in range(1, ws.max_column + 1):
            max_length = 0
            for row in range(1, ws.max_row + 1):
                cell_value = str(ws.cell(row=row, column=col).value or "")
                if len(cell_value) > max_length:
                    max_length = len(cell_value)
            ws.column_dimensions[get_column_letter(col)].width = min(max_length + 2, 50)

        return ws

    def _save(self, filename):
        """Save the workbook"""
        output_path = os.path.join(self.base_dir, filename)
        try:
            self.wb.save(output_path)
            print(f"Excel file saved: {output_path}")
        except Exception as e:
            print(f"Error saving Excel: {e}")


class MtnCsvToXlsx(BaseCsvToXlsx):
    def generate(self):
        files = os.listdir(self.base_dir) if os.path.exists(self.base_dir) else []
        for f in sorted(files):
            if f.startswith("Requisition_IdNumero_") and f.endswith(".csv"):
                self._create_sheet("Abonne", f,
                    ["Numero", "Nom et Prenom", "Date Naissance", "Numero CNI",
                     "Date Exp CNI", "Quartier", "Nationalite"])
            elif f.startswith("Requisition_Listing_") and f.endswith(".csv"):
                self._create_sheet("Listing", f,
                    ["NumeroAppelant", "LocalisationNumeroAppelant", "IMEINumeroAppelant",
                     "DateDebutAppel", "DureeAppel", "NumeroAppele"])
            elif f.startswith("IdentificationAbonnees_") and f.endswith(".csv"):
                self._create_sheet("Identification des abonnes", f,
                    ["Numero", "Nom et Prenom", "Date Naissance", "Numero CNI",
                     "Date Exp CNI", "Quartier", "Nationalite"])
            elif f.startswith("Statistiques_Localisation_") and f.endswith(".csv"):
                self._create_sheet("Statistique Localisation", f,
                    ["Localisation", "Occurence"])
            elif f.startswith("Statistiques_") and f.endswith(".csv"):
                self._create_sheet("Statistique Appel", f,
                    ["NumeroAppelant", "Occurence", "DureeAppel"])
        # Remove default empty sheet if exists
        if 'Sheet' in self.wb.sheetnames and len(self.wb.sheetnames) > 1:
            del self.wb['Sheet']
        self._save(f"{self.phone}.xlsx")


class MtnCsvToXlsxMultiple(BaseCsvToXlsx):
    def generate(self):
        files = os.listdir(self.base_dir) if os.path.exists(self.base_dir) else []
        for f in sorted(files):
            if f.startswith("Requisition_IdNumero_") and f.endswith(".csv"):
                self._create_sheet("Abonne", f)
            elif f.startswith("Requisition_Listing_") and f.endswith(".csv"):
                self._create_sheet("Listing", f)
            elif f.startswith("IdentificationAbonnees_") and f.endswith(".csv"):
                self._create_sheet("Identification des abonnes", f)
            elif f.startswith("Statistiques_Localisation_") and f.endswith(".csv"):
                self._create_sheet("Statistique Localisation", f)
            elif f.startswith("Statistiques_") and f.endswith(".csv"):
                self._create_sheet("Statistique Appel", f)
            elif f.startswith("frequenceCellule_") and f.endswith(".csv"):
                self._create_sheet("Frequence par cellule", f)
            elif f.startswith("frequenceCorrespondance_") and f.endswith(".csv"):
                self._create_sheet("Frequence Correspondant", f)
            elif f.startswith("frequenceDureeAppel_") and f.endswith(".csv"):
                self._create_sheet("Frequence par Duree appel", f)
            elif f.startswith("frequenceImei_") and f.endswith(".csv"):
                self._create_sheet("Frequence par IMEI", f)
            elif f.startswith("sharedImei_") and f.endswith(".csv"):
                self._create_sheet("Shared IMEI", f)
        if 'Sheet' in self.wb.sheetnames and len(self.wb.sheetnames) > 1:
            del self.wb['Sheet']

        # Set document properties
        self.wb.properties.creator = self.demandeur_requisition
        self.wb.properties.description = f"Requisition {self.phone}"
        self.wb.properties.title = f"CDR {self.phone}"
        self.wb.properties.category = "CDR"
        self._save(f"{self.phone}.xlsx")


class OrangeCsvToXlsx(BaseCsvToXlsx):
    def generate(self):
        files = os.listdir(self.base_dir) if os.path.exists(self.base_dir) else []
        for f in sorted(files):
            if f.startswith("Requisition_IdNumero_") and f.endswith(".csv"):
                self._create_sheet("Abonne", f,
                    ["Numero", "Nom et Prenom", "Date Naissance", "Numero CNI",
                     "Date Exp CNI", "Quartier"])
            elif f.startswith("Requisition_Listing_Emis_") and f.endswith(".csv"):
                self._create_sheet("Listing Appel", f,
                    ["NumeroAppelant", "LocalisationNumeroAppelant", "IMEINumeroAppelant",
                     "DateDebutAppel", "DureeAppel", "NumeroAppele"])
            elif f.startswith("Requisition_Listing_SMS_") and f.endswith(".csv"):
                self._create_sheet("Listing SMS", f,
                    ["NumeroEnvoi", "LocalisationNumeroDest", "IMEINumeroDest",
                     "DateSMS", "NumeroDest"])
            elif f.startswith("IdentificationAbonnees_") and f.endswith(".csv"):
                self._create_sheet("Identification des abonnes", f,
                    ["Numero", "Nom et Prenom", "Date Naissance", "Numero CNI",
                     "Date Exp CNI", "Quartier"])
            elif f.startswith("Statistiques_Localisation_") and f.endswith(".csv"):
                self._create_sheet("Statistique Localisation", f,
                    ["Localisation", "Occurence"])
            elif f.startswith("Statistiques_") and f.endswith(".csv"):
                self._create_sheet("Statistique Appel", f,
                    ["NumeroAppelant", "Occurence", "DureeAppel"])
        if 'Sheet' in self.wb.sheetnames and len(self.wb.sheetnames) > 1:
            del self.wb['Sheet']
        self._save(f"{self.phone}.xlsx")


class OrangeCsvToXlsxMultiple(BaseCsvToXlsx):
    def generate(self):
        files = os.listdir(self.base_dir) if os.path.exists(self.base_dir) else []
        for f in sorted(files):
            if f.startswith("Requisition_IdNumero_") and f.endswith(".csv"):
                self._create_sheet("Abonne", f)
            elif f.startswith("Requisition_Listing_SMS_") and f.endswith(".csv"):
                self._create_sheet("Listing SMS", f)
            elif f.startswith("Requisition_Listing_Emis_") and f.endswith(".csv"):
                self._create_sheet("Listing Appel", f)
            elif f.startswith("IdentificationAbonnees_") and f.endswith(".csv"):
                self._create_sheet("Identification des abonnes", f)
            elif f.startswith("Statistiques_Localisation_") and f.endswith(".csv"):
                self._create_sheet("Statistique Localisation", f)
            elif f.startswith("Statistiques_") and f.endswith(".csv"):
                self._create_sheet("Statistique Appel", f)
            elif f.startswith("frequenceCellule_") and f.endswith(".csv"):
                self._create_sheet("Frequence par cellule", f)
            elif f.startswith("frequenceCorrespondance_") and f.endswith(".csv"):
                self._create_sheet("Frequence Correspondant", f)
            elif f.startswith("frequenceDureeAppel_") and f.endswith(".csv"):
                self._create_sheet("Frequence par Duree appel", f)
            elif f.startswith("frequenceImei_") and f.endswith(".csv"):
                self._create_sheet("Frequence par IMEI", f)
            elif f.startswith("sharedImei_") and f.endswith(".csv"):
                self._create_sheet("Shared IMEI", f)
        if 'Sheet' in self.wb.sheetnames and len(self.wb.sheetnames) > 1:
            del self.wb['Sheet']
        self.wb.properties.creator = self.demandeur_requisition
        self._save(f"{self.phone}.xlsx")


class NexttelCsvToXlsx(BaseCsvToXlsx):
    def generate(self):
        files = os.listdir(self.base_dir) if os.path.exists(self.base_dir) else []
        for f in sorted(files):
            if f.startswith("Requisition_Identification_Numero_") and f.endswith(".csv"):
                self._create_sheet("Abonne", f,
                    ["Numero", "Nom et Prenom", "Date Naissance", "Numero CNI", "Date Exp CNI"])
            elif f.startswith("Requisition_Listing_") and f.endswith(".csv"):
                self._create_sheet("Listing", f,
                    ["NumeroAppelant", "LocalisationNumeroAppelant", "IMEINumeroAppelant",
                     "DateDebutAppel", "DureeAppel", "NumeroAppele"])
            elif f.startswith("IdentificationAbonnees_") and f.endswith(".csv"):
                self._create_sheet("Identification des abonnes", f,
                    ["Numero", "Nom et Prenom", "Date Naissance", "Numero CNI", "Date Exp CNI"])
            elif f.startswith("Statistiques_Localisation_") and f.endswith(".csv"):
                self._create_sheet("Statistique Localisation", f,
                    ["Localisation", "Occurence"])
            elif f.startswith("Statistiques_") and f.endswith(".csv"):
                self._create_sheet("Statistique Appel", f,
                    ["NumeroAppelant", "Occurence", "DureeAppel"])
        if 'Sheet' in self.wb.sheetnames and len(self.wb.sheetnames) > 1:
            del self.wb['Sheet']
        self._save(f"{self.phone}.xlsx")


class NexttelCsvToXlsxMultiple(BaseCsvToXlsx):
    def generate(self):
        files = os.listdir(self.base_dir) if os.path.exists(self.base_dir) else []
        for f in sorted(files):
            if f.startswith("Requisition_Identification_Numero_") and f.endswith(".csv"):
                self._create_sheet("Abonne", f)
            elif f.startswith("Requisition_Listing_") and f.endswith(".csv"):
                self._create_sheet("Listing", f)
            elif f.startswith("IdentificationAbonnees_") and f.endswith(".csv"):
                self._create_sheet("Identification des abonnes", f)
            elif f.startswith("Statistiques_Localisation_") and f.endswith(".csv"):
                self._create_sheet("Statistique Localisation", f)
            elif f.startswith("Statistiques_") and f.endswith(".csv"):
                self._create_sheet("Statistique Appel", f)
            elif f.startswith("frequenceCellule_") and f.endswith(".csv"):
                self._create_sheet("Frequence par cellule", f)
            elif f.startswith("frequenceCorrespondance_") and f.endswith(".csv"):
                self._create_sheet("Frequence Correspondant", f)
            elif f.startswith("frequenceDureeAppel_") and f.endswith(".csv"):
                self._create_sheet("Frequence par Duree appel", f)
            elif f.startswith("frequenceImei_") and f.endswith(".csv"):
                self._create_sheet("Frequence par IMEI", f)
            elif f.startswith("sharedImei_") and f.endswith(".csv"):
                self._create_sheet("Shared IMEI", f)
        if 'Sheet' in self.wb.sheetnames and len(self.wb.sheetnames) > 1:
            del self.wb['Sheet']
        self.wb.properties.creator = self.demandeur_requisition
        self._save(f"{self.phone}.xlsx")


class MomoCsvToXlsx(BaseCsvToXlsx):
    def generate(self):
        files = os.listdir(self.base_dir) if os.path.exists(self.base_dir) else []
        for f in sorted(files):
            if f.startswith("Requisition_IdNumero_") and f.endswith(".csv"):
                self._create_sheet("Abonne", f)
            elif f.startswith("Requisition_Listing_") and f.endswith(".csv"):
                self._create_sheet("Listing", f,
                    ["IdTransaction", "DateTransaction", "IdEmetteur", "NumeroEmetteur",
                     "IdRecepteur", "NumeroRecepteur", "Montant", "LIQUIDE"])
            elif f.startswith("Statistiques_Localisation_") and f.endswith(".csv"):
                self._create_sheet("Statistique Localisation", f)
            elif f.startswith("Statistiques_") and f.endswith(".csv"):
                self._create_sheet("Statistique Appel", f)
            elif f.startswith("IdentificationAbonnees_") and f.endswith(".csv"):
                self._create_sheet("Identification des abonnes", f)
        if 'Sheet' in self.wb.sheetnames and len(self.wb.sheetnames) > 1:
            del self.wb['Sheet']
        self._save(f"{self.phone}_MOMO.xlsx")
